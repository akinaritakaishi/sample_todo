"""Generate the fictional sample files for demo 2 (expense check).

Usage: python3 docs/demo/expense/tools/generate.py
Requires: openpyxl, reportlab (not part of front/'s dependencies) and a Japanese
font (edit FONT_PATH if WenQuanYi Zen Hei is not installed).
All people, shops and registration numbers are fictional.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib.pagesizes import A5, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

BASE = Path(__file__).resolve().parent.parent
# A Japanese TTF/TTC font embedded into the receipts so any PDF viewer can render them.
FONT_PATH = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
FONT = "JP"
FORM_FIELDS = ["申請日", "申請者", "所属", "利用日", "区分", "内容", "支払先", "金額", "参加人数", "相手先", "目的", "承認者"]
LEDGER_HEADERS = ["No", "フォルダ名", "申請者", "利用日", "区分", "支払先", "金額", "人数", "1人あたり", "判定", "理由（条番号）", "通知日時"]

# folder, form values, receipt values
APPLICATIONS = [
    ("山田太郎_タクシー_20260908",
     dict(申請日="2026-10-01", 申請者="山田太郎", 所属="営業部", 利用日="2026-09-08", 区分="交通費", 内容="タクシー代",
          支払先="みどり交通株式会社", 金額=2860, 目的="ひかり商事様へ訪問"),
     dict(amount=2860, date="2026年9月8日", payee="みどり交通株式会社", note="タクシー代として", reg="T9000000000011")),
    ("山田太郎_取引先会食_20260912",
     dict(申請日="2026-10-01", 申請者="山田太郎", 所属="営業部", 利用日="2026-09-12", 区分="交際費", 内容="取引先との会食",
          支払先="和食処 さくら亭", 金額=18000, 参加人数=4, 相手先="ひかり商事 2名", 目的="新規案件の打ち合わせ"),
     dict(amount=18000, date="2026年9月12日", payee="和食処 さくら亭", note="ご飲食代として（4名様）", reg="T9000000000022")),
    ("山田太郎_取引先会食_20260918",
     dict(申請日="2026-10-02", 申請者="山田太郎", 所属="営業部", 利用日="2026-09-18", 区分="交際費", 内容="取引先との会食",
          支払先="ビストロ ほしぞら", 金額=16500, 参加人数=3, 相手先="あおば物産 2名", 目的="契約更新のお礼"),
     dict(amount=16500, date="2026年9月18日", payee="ビストロ ほしぞら", note="ご飲食代として（3名様）", reg="T9000000000033")),
    ("佐藤花子_書籍購入_20260915",
     dict(申請日="2026-10-02", 申請者="佐藤花子", 所属="総務部", 利用日="2026-09-15", 区分="物品購入", 内容="書籍（労務管理の実務）",
          支払先="こもれび書店", 金額=3300, 目的="業務知識の習得"),
     dict(amount=3300, date="2026年9月15日", payee="こもれび書店", note="書籍代として", reg="T9000000000044")),
    ("佐藤花子_備品購入_20260922",
     dict(申請日="2026-10-03", 申請者="佐藤花子", 所属="総務部", 利用日="2026-09-22", 区分="物品購入", 内容="ファイルボックス・ラベルシール",
          支払先="文具のまるやま", 金額=4980, 目的="書類整理用の備品"),
     dict(amount=4980, date="2026年9月22日", payee="文具のまるやま", note="文具代として", reg=None)),
    ("鈴木一郎_出張交通費_20260925",
     dict(申請日="2026-10-05", 申請者="鈴木一郎", 所属="開発部", 利用日="2026-09-25", 区分="交通費", 内容="新幹線（東京⇔新大阪 往復）",
          支払先="ひまわり旅行センター", 金額=29440, 目的="大阪営業所での導入支援"),
     dict(amount=27440, date="2026年9月25日", payee="ひまわり旅行センター", note="乗車券・特急券代として", reg="T9000000000055")),
]

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def build_form(values=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "申請書"
    ws["A1"] = "経費精算申請書（株式会社サンプル研修）"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "※勉強会デモ用の架空の書式です"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 40
    for i, field in enumerate(FORM_FIELDS, start=4):
        ws.cell(row=i, column=1, value=field).font = HEADER_FONT
        ws.cell(row=i, column=1).fill = HEADER_FILL
        if values and field in values:
            ws.cell(row=i, column=2, value=values[field])
    return wb


def build_receipt(path, r):
    c = canvas.Canvas(str(path), pagesize=landscape(A5))
    w, h = landscape(A5)
    c.setFont(FONT, 22)
    c.drawCentredString(w / 2, h - 60, "領 収 書")
    c.setFont(FONT, 13)
    c.drawString(50, h - 110, "株式会社サンプル研修 様")
    c.line(50, h - 115, 260, h - 115)
    c.setFont(FONT, 20)
    c.drawCentredString(w / 2, h - 170, f"金額  ¥{r['amount']:,}-（税込）")
    c.setFont(FONT, 11)
    c.drawString(50, h - 210, f"但し {r['note']}")
    c.drawString(50, h - 230, "上記正に領収いたしました。")
    c.drawString(w - 250, h - 280, r["date"])
    c.drawString(w - 250, h - 300, r["payee"])
    if r["reg"]:
        c.drawString(w - 250, h - 320, f"登録番号 {r['reg']}")
    c.setFont(FONT, 9)
    c.drawString(50, 30, "※デモ用の架空の領収書です。実在の店舗・事業者とは関係ありません。")
    c.save()


def build_ledger(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "10月申請分"
    widths = [5, 32, 10, 12, 10, 22, 10, 6, 10, 10, 50, 18]
    for col, (header, width) in enumerate(zip(LEDGER_HEADERS, widths), start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"OK,要確認,差し戻し"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("J2:J1000")
    wb.save(path)


def main():
    pdfmetrics.registerFont(TTFont(FONT, FONT_PATH, subfontIndex=0))
    rules = BASE / "経費精算ルール"
    rules.mkdir(exist_ok=True)
    build_form().save(rules / "申請書テンプレート.xlsx")
    for folder, form, receipt in APPLICATIONS:
        d = BASE / "10月申請分" / folder
        d.mkdir(parents=True, exist_ok=True)
        build_form(form).save(d / "申請書.xlsx")
        build_receipt(d / "領収書.pdf", receipt)
    ledger_dir = BASE / "経理用"
    ledger_dir.mkdir(exist_ok=True)
    build_ledger(ledger_dir / "10月申請分_経費台帳.xlsx")


if __name__ == "__main__":
    main()
